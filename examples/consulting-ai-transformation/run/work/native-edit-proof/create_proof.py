"""Edit an existing assembled deck copy and verify the Office reader's result."""
import argparse
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import cv2
import numpy as np
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.chart.data import CategoryChartData


HERE = Path(__file__).resolve().parent
SOURCE = HERE.parents[1]
ROOT = HERE.parents[4]
RUNTIME = ROOT / "slidepoise/scripts/slidepoise_runtime.py"
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write(path, data):
    path.write_text(json.dumps(data, indent=2) + "\n")


def orange_bar_pixels(preview, chart_box, dimensions):
    with Image.open(preview) as image:
        x, y, width, height = chart_box
        sx, sy = image.width / dimensions[0], image.height / dimensions[1]
        crop = np.asarray(image.convert("RGB").crop((round(x*sx), round(y*sy), round((x+width)*sx), round((y+height)*sy))))
    mask = (np.abs(crop.astype(np.int16) - [253, 81, 8]).max(axis=2) <= 12).astype(np.uint8)
    _, _, statistics, _ = cv2.connectedComponentsWithStats(mask)
    component = statistics[1:][np.argmax(statistics[1:, cv2.CC_STAT_AREA])]
    return {"width_px": int(component[cv2.CC_STAT_WIDTH]), "height_px": int(component[cv2.CC_STAT_HEIGHT])}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-sha256", required=True)
    args = parser.parse_args()
    original = SOURCE / "deliverables/presentation.pptx"
    assert digest(original) == args.source_sha256, "Official deck changed before the editing demonstration"
    manifest_path = SOURCE / "work/deck-scenes.json"
    manifest = json.loads(manifest_path.read_text())
    inputs = [original, manifest_path, SOURCE / "work/resolved-config.json", SOURCE / "work/deck-design.json"]
    for entry in manifest["slides"]:
        inputs.append((manifest_path.parent / entry["scene"]).resolve())
        work = SOURCE / "slides" / entry["slide_id"] / "work"
        inputs.extend(work / name for name in ("accepted-slide.png", "semantic-map.json", "reconstruction-handoff.json"))
    source_bindings = {str(path.relative_to(SOURCE)): digest(path) for path in inputs}
    write(HERE / "source-bindings.json", source_bindings)
    original_copy = HERE / "official-copy.pptx"
    edited = HERE / "edited-copy.pptx"
    if edited.exists() or (HERE / "edited-preview").exists():
        raise SystemExit("This proof already exists. Preserve its evidence before making another demonstration.")
    shutil.copy2(original, original_copy)
    shutil.copy2(original, edited)
    deck = Presentation(edited)
    slide = deck.slides[1]
    title = next(shape for shape in slide.shapes if shape.name == "headline")
    old_title = title.text
    assert len(title.text_frame.paragraphs) == len(title.text_frame.paragraphs[0].runs) == 1
    new_title = "Revise the analysis without rebuilding the slide"
    title.text_frame.paragraphs[0].runs[0].text = new_title
    chart_shape = next(shape for shape in slide.shapes if shape.has_chart and shape.name == "annual-capacity-chart")
    chart = chart_shape.chart
    before_values = list(chart.series[0].values)
    assert before_values == [1944, 1440, 1080]
    data = CategoryChartData()
    data.categories = [category.label for category in chart.plots[0].categories]
    data.add_series(chart.series[0].name, [1620, 1440, 1080])
    chart.replace_data(data)
    chart_part = str(chart.part.partname).lstrip("/")
    workbook_part = str(chart.part.chart_workbook.xlsx_part.partname).lstrip("/")
    deck.save(edited)
    with ZipFile(edited) as package:
        chart_xml = ET.fromstring(package.read(chart_part))
        cache = [float(node.text) for node in chart_xml.findall(".//c:val/c:numRef/c:numCache/c:pt/c:v", NS)]
        assert cache == [1620, 1440, 1080]
        with ZipFile(io.BytesIO(package.read(workbook_part))) as workbook:
            sheet = ET.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
            workbook_value = float(sheet.find(".//s:c[@r='B2']/s:v", NS).text)
        assert workbook_value == 1620
        assert chart_xml.find(".//c:externalData", NS) is not None
        labels = chart_xml.findall(".//c:dLbls", NS)
        assert labels and all(label.find("c:showVal", NS).get("val") == "1" for label in labels)
        assert all(label.find("c:txPr/a:bodyPr", NS).get("wrap") == "none" for label in labels)
        assert all(label.find("c:tx", NS) is None for group in labels for label in group.findall("c:dLbl", NS))
        assert all(group.find("c:numFmt", NS).get("formatCode") == '#,##0" h"' for group in labels)
    commands = []

    def run(*command):
        result = subprocess.run([str(arg) for arg in command], cwd=ROOT, text=True, capture_output=True,
                                env={**os.environ, "PYTHON": sys.executable, "PYTHONDONTWRITEBYTECODE": "1"}, timeout=120)
        commands.append({"command": [str(arg) for arg in command], "exit_code": result.returncode,
                         "stdout": result.stdout, "stderr": result.stderr})
        write(HERE / "commands.json", commands)
        result.check_returncode()
        return result.stdout

    for name, path in (("official", original_copy), ("edited", edited)):
        run(sys.executable, RUNTIME, "render-deck-preview", "--pptx", path,
            "--output-dir", HERE / f"{name}-preview", "--dpi", "120")
    pages = run("pdftotext", "-layout", HERE / "edited-preview/presentation.pdf", "-").split("\f")
    write(HERE / "actual-pdf-text.json", {"page_number": 2, "text": pages[1]})
    assert new_title in pages[1]
    assert any(re.search(r"1,620\s+h\b", line) for line in pages[1].splitlines())
    assert "1,944" not in pages[1]
    comparison = HERE / "comparison.png"
    run(sys.executable, ROOT / "slidepoise/scripts/make_visual_comparison.py",
        "--left", HERE / "official-preview/slide-002.png", "--right", HERE / "edited-preview/slide-002.png",
        "--left-label", "Official planning scenario", "--right-label", "Native edit demonstration only", "--output", comparison)
    untouched_pages = []
    for page in (1, 3, 4, 5):
        name = f"slide-{page:03d}.png"
        with Image.open(HERE / "official-preview" / name) as first, Image.open(HERE / "edited-preview" / name) as second:
            equal = first.size == second.size and ImageChops.difference(first.convert("RGB"), second.convert("RGB")).getbbox() is None
        untouched_pages.append({"page_number": page, "pixels_identical": equal})
    assert all(item["pixels_identical"] for item in untouched_pages)
    entry = manifest["slides"][1]
    scene = json.loads((manifest_path.parent / entry["scene"]).resolve().read_text())
    chart_box = next(obj["bbox_px"] for obj in scene["objects"] if obj["id"] == chart_shape.name)
    before_bar = orange_bar_pixels(HERE / "official-preview/slide-002.png", chart_box, scene["dimensions_px"])
    after_bar = orange_bar_pixels(HERE / "edited-preview/slide-002.png", chart_box, scene["dimensions_px"])
    assert abs(after_bar["width_px"] - before_bar["width_px"] * 1620 / 1944) <= 2
    write(HERE / "bar-edit-evidence.json", {"chart_bbox_px": chart_box, "dimensions_px": scene["dimensions_px"],
          "official_bar": before_bar, "edited_bar": after_bar, "expected_value_ratio": 1620 / 1944,
          "official_render_sha256": digest(HERE / "official-preview/slide-002.png"),
          "edited_render_sha256": digest(HERE / "edited-preview/slide-002.png")})
    assert all(digest(SOURCE / path) == sha for path, sha in source_bindings.items())
    write(HERE / "edit-proof.json", {
        "schema_version": "1.0", "method": "Edited the existing native textbox and chart in a copy of the assembled runtime output using python-pptx. No slide was rebuilt.",
        "official_source": "../../deliverables/presentation.pptx", "official_source_sha256": args.source_sha256,
        "source": original_copy.name, "source_sha256": digest(original_copy),
        "edited_pptx": edited.name, "edited_pptx_sha256": digest(edited), "page_number": 2,
        "title_edit": {"object_name": title.name, "before": old_title, "after": new_title},
        "chart_edit": {"object_name": chart_shape.name, "before": before_values, "after": cache,
                       "embedded_workbook_first_value": workbook_value, "native_labels_bound_to_values": True,
                       "native_wrap_policy": "none", "actual_pdf_label": "1,620 h"},
        "source_preserved": True, "editable_numeric_textboxes_created": 0, "unmodified_pages": untouched_pages,
        "edited_preview": "edited-preview/slide-002.png", "edited_preview_sha256": digest(HERE / "edited-preview/slide-002.png"),
        "comparison": comparison.name, "comparison_sha256": digest(comparison),
        "disclosure": "This separate copy demonstrates native editing only. Changing 1,944 to 1,620 does not revise the official planning scenario or its calculations. The official deck, scenes, semantic maps, notes and generated targets remain unchanged. The assumption table is an independent business view and is intentionally unchanged in this editing demonstration."
    })
    print(json.dumps({"official_source_preserved": True, "native_workbook_and_pdf_value": workbook_value,
                      "other_pages_pixels_identical": len(untouched_pages), "comparison": str(comparison)}, indent=2))


if __name__ == "__main__":
    main()
